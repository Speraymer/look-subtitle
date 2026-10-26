"""PPT 画面：下载平台抽的图、裁掉浏览器边框、拼成 Word 讲义。

两条图源：
  1) 平台 AI 抽帧（一定有，但有一部分是投影幕布的相机画面，糊）
  2) 外部课件（QQ 群里老师发的 PDF/PPT 导出的图片，最清楚）
外部课件放在 works/<id>/ppt/slides/ 下，用 `look slides` 登记，优先使用。
"""

import json
import os
import re
import subprocess
from pathlib import Path

from PIL import Image

from .paths import Work

CROP_TOP, CROP_BOTTOM = 95, 28   # 全屏截图里浏览器界面占的高度


def mmss(t):
    t = int(round(t))
    return f"{t // 60:02d}:{t % 60:02d}"


# ---------- 图 ----------

def trim_sides(im):
    """去掉 PDF 阅读器两侧的灰底，只留页面本身。"""
    import numpy as np

    g = np.array(im.convert("L"))
    white = (g > 252).mean(axis=0)
    cols = np.where(white > 0.10)[0]
    if len(cols) < 10:
        return im
    x0, x1 = max(0, cols[0] - 3), min(im.width, cols[-1] + 3)
    if x1 - x0 < im.width * 0.4:
        return im
    return im.crop((x0, 0, x1, im.height))


def download_pages(w: Work, session=None, force=False):
    """从 pages.json 里把图拉下来。"""
    pages = json.load(open(w.ppt_pages, encoding="utf-8"))
    w.ensure_dirs()
    if session is None:
        import requests
        import urllib3

        urllib3.disable_warnings()
        session = requests.Session()
        session.verify = False
        session.headers.update({"User-Agent": "Mozilla/5.0"})
    ok = 0
    for i, p in enumerate(pages):
        dest = w.ppt_raw / f"{i:02d}.jpg"
        if dest.exists() and not force:
            ok += 1
            continue
        try:
            r = session.get(p["img"], timeout=30)
            if r.status_code == 200:
                dest.write_bytes(r.content)
                ok += 1
        except Exception as e:
            print(f"    {i:02d} 失败 {e}")
    print(f"  PPT 图 {ok}/{len(pages)} 张 -> {w.ppt_raw.name}/")
    return pages


def crop_pages(w: Work):
    """裁浏览器边框，输出到 ppt/crop/。返回 [(path, 是否小图)]。"""
    w.ensure_dirs()
    out = []
    for i in range(1000):
        src = w.ppt_raw / f"{i:02d}.jpg"
        if not src.exists():
            break
        im = Image.open(src).convert("RGB")
        small = im.width < 900
        if not small:
            im = im.crop((0, CROP_TOP, im.width, im.height - CROP_BOTTOM))
            im = trim_sides(im)
        dest = w.ppt_crop / f"{i:02d}.jpg"
        im.save(dest, quality=88)
        out.append((dest, small))
    return out


def _image_files(src):
    return sorted(
        [p for p in src.iterdir() if p.suffix.lower() in (".jpg", ".jpeg", ".png", ".webp")],
        key=lambda p: (len(p.stem), p.stem),
    )


def _render_courseware(source, dest):
    """把 PDF/PPT 渲染成图片；图片目录直接返回。"""
    source = Path(source).resolve()
    if source.is_dir():
        return _image_files(source)
    dest.mkdir(parents=True, exist_ok=True)
    if source.suffix.lower() == ".pdf":
        try:
            import pymupdf as fitz
        except ImportError as e:
            raise SystemExit("渲染 PDF 需要 PyMuPDF：python -m pip install pymupdf") from e
        with fitz.open(source) as doc:
            for i, page in enumerate(doc):
                page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False).save(dest / f"{i:04d}.jpg")
        return _image_files(dest)
    if source.suffix.lower() not in (".ppt", ".pptx", ".pptm"):
        raise SystemExit(f"不支持的课件格式：{source.suffix}")

    env = os.environ.copy()
    env.update({"LOOK_SLIDE_SOURCE": str(source), "LOOK_SLIDE_DEST": str(dest)})
    script = (
        "$ErrorActionPreference='Stop';"
        "$app=New-Object -ComObject PowerPoint.Application;"
        "try{$p=$app.Presentations.Open($env:LOOK_SLIDE_SOURCE,$true,$true,$false);"
        "$p.Export($env:LOOK_SLIDE_DEST,'JPG',1920,1080);$p.Close()}finally{$app.Quit()}"
    )
    try:
        subprocess.run(
            ["powershell", "-NoProfile", "-NonInteractive", "-Command", script],
            check=True, env=env, capture_output=True, text=True,
        )
    except subprocess.CalledProcessError as e:
        raise SystemExit("PowerPoint 渲染失败：" + (e.stderr or e.stdout or str(e)).strip()) from e
    return _image_files(dest)


def _hash_variants(path):
    """多种中心裁剪的 dHash；兼容全屏截图和拍幕布的相机画面。"""
    import numpy as np

    im = Image.open(path).convert("L")
    boxes = ((0, 0, 1, 1), (.05, .05, .95, .95), (.12, .10, .88, .90))
    out = []
    for x0, y0, x1, y1 in boxes:
        crop = im.crop((int(im.width*x0), int(im.height*y0), int(im.width*x1), int(im.height*y1)))
        a = np.asarray(crop.resize((33, 18)))
        out.append(a[:, 1:] > a[:, :-1])
    return out


def _distance(left, right):
    import numpy as np

    return min(float(np.mean(a != b)) for a in left for b in right)


def _features(path):
    """ORB 局部特征能认出被浏览器边框或相机裁切过的同一页。"""
    import numpy as np
    from skimage.color import rgb2gray
    from skimage.feature import ORB
    from skimage.transform import resize

    a = np.asarray(Image.open(path).convert("RGB"), dtype=np.float32) / 255.0
    a = resize(a, (max(1, int(a.shape[0] * 600 / a.shape[1])), 600), anti_aliasing=True)
    orb = ORB(n_keypoints=400, fast_threshold=.03)
    try:
        orb.detect_and_extract(rgb2gray(a))
        return orb.keypoints, orb.descriptors
    except RuntimeError:
        return np.empty((0, 2)), np.empty((0, 256), dtype=bool)


def _sift_features(path):
    import numpy as np
    from skimage.color import rgb2gray
    from skimage.feature import SIFT
    from skimage.transform import resize

    a = np.asarray(Image.open(path).convert("RGB"), dtype=np.float32) / 255.0
    a = resize(a, (max(1, int(a.shape[0] * 600 / a.shape[1])), 600), anti_aliasing=True)
    sift = SIFT(n_octaves=5, n_scales=3)
    try:
        sift.detect_and_extract(rgb2gray(a))
        return sift.keypoints, sift.descriptors
    except RuntimeError:
        return np.empty((0, 2)), np.empty((0, 128), dtype=np.float32)


def _feature_scores(frames, slides, fallback):
    import numpy as np
    from skimage.feature import match_descriptors
    from skimage.measure import ransac
    from skimage.transform import ProjectiveTransform

    ff = [_features(p) for p in frames]
    sf = [_features(p) for p in slides]
    score = fallback.copy()
    inliers = np.zeros_like(score, dtype=np.int16)
    for i, (fk, fd) in enumerate(ff):
        candidates = []
        for j, (sk, sd) in enumerate(sf):
            if not len(fd) or not len(sd):
                continue
            matches = match_descriptors(fd, sd, metric="hamming", cross_check=True, max_ratio=.9)
            candidates.append((len(matches), j, matches))
        # shortcut: only the ten strongest descriptor candidates get RANSAC; widen this if similar decks miss pages.
        for _, j, matches in sorted(candidates, reverse=True)[:10]:
            if len(matches) < 8:
                continue
            try:
                _, mask = ransac(
                    (sf[j][0][matches[:, 1]], fk[matches[:, 0]]),
                    ProjectiveTransform, min_samples=4, residual_threshold=3,
                    max_trials=80, rng=0,
                )
            except (ValueError, RuntimeError):
                continue
            count = int(mask.sum()) if mask is not None else 0
            inliers[i, j] = count
            if count >= 6:
                score[i, j] = .85 * max(0.0, 1.0 - count / 30.0) + .15 * fallback[i, j]

    # Camera frames are tiny, tinted and partly blocked; SIFT is slower but much steadier on those only.
    low_res = []
    for i, path in enumerate(frames):
        with Image.open(path) as im:
            if im.width < 600:
                low_res.append(i)
    if low_res:
        ssf = [_sift_features(p) for p in slides]
        for i in low_res:
            fk, fd = _sift_features(frames[i])
            candidates = []
            likely = set(np.argsort(fallback[i])[:30])
            for j in np.flatnonzero(inliers[i]):
                likely.update(range(max(0, j - 2), min(len(slides), j + 3)))
            for j in sorted(likely):
                sk, sd = ssf[j]
                if not len(fd) or not len(sd):
                    continue
                matches = match_descriptors(fd, sd, metric="euclidean", cross_check=True, max_ratio=.8)
                candidates.append((len(matches), j, matches))
            # shortcut: only the ten strongest SIFT candidates get RANSAC; widen this if camera frames miss pages.
            for _, j, matches in sorted(candidates, reverse=True)[:10]:
                if len(matches) < 8:
                    continue
                try:
                    _, mask = ransac(
                        (ssf[j][0][matches[:, 1]], fk[matches[:, 0]]),
                        ProjectiveTransform, min_samples=4, residual_threshold=4,
                        max_trials=120, rng=0,
                    )
                except (ValueError, RuntimeError):
                    continue
                count = int(mask.sum()) if mask is not None else 0
                if count > inliers[i, j]:
                    inliers[i, j] = count
                    score[i, j] = .85 * max(0.0, 1.0 - count / 30.0) + .15 * fallback[i, j]
    return score, inliers


def _monotone_match(frames, slides):
    """找帧到课件页的单调最佳路径，避免页数不同时按比例错配。"""
    import numpy as np

    fh = [_hash_variants(p) for p in frames]
    sh = [_hash_variants(p) for p in slides]
    fallback = np.array([[_distance(a, b) for b in sh] for a in fh], dtype=np.float32)
    score, inliers = _feature_scores(frames, slides, fallback)
    dp = score[0].copy()
    parents = []
    jump = .00005
    for i in range(1, len(frames)):
        best, best_at = float("inf"), 0
        cur, parent = np.empty(len(slides)), np.empty(len(slides), dtype=np.int32)
        for j in range(len(slides)):
            candidate = float(dp[j]) - jump*j
            if candidate < best:
                best, best_at = candidate, j
            cur[j] = score[i, j] + best + jump*j
            parent[j] = best_at
        dp, parents = cur, parents + [parent]
    path = [int(np.argmin(dp))]
    for parent in reversed(parents):
        path.append(int(parent[path[-1]]))
    path.reverse()

    # A whole lecture collapsing onto one template page is a false match, not a useful alignment.
    enough_sequence = len(set(path)) >= min(len(frames), max(2, (len(frames) + 4) // 5))

    rows = []
    for i, path_j in enumerate(path):
        feature_j = int(np.argmax(inliers[i]))
        ranked_inliers = np.sort(inliers[i])
        best_inliers = int(ranked_inliers[-1])
        next_inliers = int(ranked_inliers[-2]) if len(ranked_inliers) > 1 else 0
        feature_ok = best_inliers >= 25 or (best_inliers >= 12 and best_inliers - next_inliers >= 5)
        j = feature_j if feature_ok else path_j
        ordered = np.sort(score[i])
        margin = float(ordered[1] - score[i, j]) if len(ordered) > 1 else 1.0
        distance = float(score[i, j])
        hash_ok = not inliers[i, j] and (distance <= .16 or (distance <= .26 and margin >= .025))
        confident = bool(feature_ok or (enough_sequence and hash_ok))
        rows.append({"frame": i, "slide": j, "distance": round(distance, 4),
                     "margin": round(margin, 4), "inliers": int(inliers[i, j]),
                     "confident": confident})
    return rows


def use_external_slides(w: Work, source, map_file=None):
    """按图像内容单调匹配高清课件；不确定的帧保留平台原图。"""
    src = Path(source)
    render_dir = w.ppt_dir / ("slides-" + re.sub(r"[^0-9A-Za-z_-]+", "-", src.stem).strip("-")[:40])
    files = _render_courseware(src, render_dir)
    files = sorted(
        files,
        key=lambda p: (len(p.stem), p.stem),
    )
    if not files:
        raise SystemExit(f"{src} 里没有图片")

    pages = json.load(open(w.ppt_pages, encoding="utf-8")) if w.ppt_pages.exists() else []
    if any(w.ppt_raw.glob("*.jpg")):
        crop_pages(w)
    frames = sorted(w.ppt_crop.glob("*.jpg"), key=lambda p: (len(p.stem), p.stem))
    if not frames:
        raise SystemExit("没有平台画面可供匹配，先跑 `look ppt`")

    if map_file:
        manual = {int(k): int(v) for k, v in json.load(open(map_file, encoding="utf-8")).items()}
        rows = [{"frame": i, "slide": j, "distance": None, "margin": None, "confident": True}
                for i, j in sorted(manual.items())]
    else:
        rows = _monotone_match(frames, files)

    w.ensure_dirs()
    copied = 0
    for row in rows:
        i, j = row["frame"], row["slide"]
        if not row["confident"]:
            continue
        if not (0 <= j < len(files)):
            continue
        dst = w.ppt_crop / f"{i:02d}.jpg"
        Image.open(files[j]).convert("RGB").save(dst, quality=90)
        copied += 1
    report = {"source": str(src.resolve()), "slides": len(files), "frames": len(frames), "matches": rows}
    with open(w.ppt_mapping, "w", encoding="utf-8") as f:
        json.dump(report, f, ensure_ascii=False, indent=1)
    print(f"  高清课件 {len(files)} 页，可信替换 {copied}/{len(frames)} 张 -> {w.ppt_crop}")
    print(f"  对齐报告 -> {w.ppt_mapping.name}")
    return copied


# ---------- Word ----------

def _set_font(run, name="宋体", size=None, bold=None, color=None):
    from docx.oxml.ns import qn
    from docx.shared import Pt

    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color:
        run.font.color.rgb = color


def _split_paragraphs(text, max_len=110):
    parts = re.split(r"(?<=[。！？])", text)
    paras, cur = [], ""
    for s in parts:
        if not s:
            continue
        cur += s
        if len(cur) >= max_len:
            paras.append(cur)
            cur = ""
    if cur:
        paras.append(cur)
    return paras


def _clean(text):
    return re.sub(r"\s+", "", text.replace("。。", "。").replace("，，", "，"))


def build_docx(w: Work, title=None, subtitle=None, duration=None, note=None):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor

    meta = json.load(open(w.meta, encoding="utf-8")) if w.meta.exists() else {}
    seg = json.load(open(w.segments, encoding="utf-8"))
    segs = seg["segments"]
    duration = float(duration or seg.get("duration") or meta.get("duration") or 0)

    pages = json.load(open(w.ppt_pages, encoding="utf-8")) if w.ppt_pages.exists() else []
    crops = crop_pages(w) if not any(w.ppt_crop.glob("*.jpg")) else [
        (p, Image.open(p).width < 600) for p in sorted(w.ppt_crop.glob("*.jpg"))
    ]

    n = max(len(crops), len(pages))
    if n == 0:
        raise SystemExit("没有画面也没有字幕，先跑 fetch/transcribe")

    bounds = [p["t"] / 1000.0 for p in pages] + [duration]
    while len(bounds) < n + 1:
        bounds.append(duration)

    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(1.8)
        s.bottom_margin = Cm(1.8)
        s.left_margin = Cm(2.0)
        s.right_margin = Cm(2.0)

    head = doc.add_paragraph()
    _set_font(head.add_run(title or meta.get("title") or f"课程讲义 · {w.sub_id}"),
              "微软雅黑", 20, True, RGBColor(0x1F, 0x3A, 0x5F))
    if subtitle:
        p = doc.add_paragraph()
        _set_font(p.add_run(subtitle), "微软雅黑", 11, False, RGBColor(0x55, 0x5F, 0x6B))

    line = f"共 {n} 个画面 · 时长 {mmss(duration)}"
    if note:
        line += " · " + note
    p = doc.add_paragraph()
    _set_font(p.add_run(line), "微软雅黑", 9, False, RGBColor(0x88, 0x8F, 0x99))
    doc.add_paragraph()

    written = 0
    for i in range(n):
        start, end = bounds[i], bounds[i + 1] if i + 1 < len(bounds) else duration
        body = "".join(_clean(s["text"]) for s in segs if start <= s["start"] < end)

        head = doc.add_paragraph()
        head.paragraph_format.space_before = Pt(10)
        head.paragraph_format.space_after = Pt(2)
        _set_font(head.add_run(f"{i+1:02d}   {mmss(start)} – {mmss(end)}"),
                  "Consolas", 10, True, RGBColor(0x2A, 0x6F, 0xD6))

        if i < len(crops):
            path, small = crops[i]
            pic = doc.add_paragraph()
            pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pic.paragraph_format.space_after = Pt(4)
            pic.add_run().add_picture(str(path), width=Cm(9.5 if small else 15.5))
            if small:
                cap = doc.add_paragraph()
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                _set_font(cap.add_run("（投影幕布的相机画面，字不清楚）"), "微软雅黑", 8, False,
                          RGBColor(0x99, 0x99, 0x99))

        for para in _split_paragraphs(body):
            q = doc.add_paragraph()
            q.paragraph_format.space_after = Pt(3)
            q.paragraph_format.line_spacing = 1.35
            _set_font(q.add_run(para), "宋体", 10.5)
        written += 1

    doc.save(w.docx)
    print(f"  讲义 {written} 节 -> {w.docx.name}（{w.docx.stat().st_size/1048576:.1f} MB）")
    return w.docx
